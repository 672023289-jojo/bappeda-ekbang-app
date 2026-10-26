import pypdf
import docx
import openpyxl
import pptx
from PIL import Image
import easyocr
import numpy as np
import os
import plotly.express as px
import sqlite3

# Menggunakan SQLite biasa sebagai Vector/Text Storage sederhana agar tidak terkena httpx.ReadTimeout
DB_RAG = "rag_storage.db"

def init_rag_db():
    conn = sqlite3.connect(DB_RAG)
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS document_chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            doc_id INTEGER,
            filename TEXT,
            opd_name TEXT,
            chunk_text TEXT
        )
    ''')
    conn.commit()
    conn.close()

init_rag_db()

def load_ocr_reader():
    return easyocr.Reader(['id', 'en'])

def extract_text_from_filepath(filepath):
    file_type = filepath.split('.')[-1].lower()
    extracted_text = ""

    try:
        if file_type == 'txt':
            with open(filepath, 'r', encoding='utf-8') as f:
                extracted_text = f.read()

        elif file_type == 'pdf':
            pdf_reader = pypdf.PdfReader(filepath)
            for page in pdf_reader.pages:
                text = page.extract_text()
                if text:
                    extracted_text += text + "\n"

        elif file_type == 'docx':
            doc = docx.Document(filepath)
            for para in doc.paragraphs:
                extracted_text += para.text + "\n"

        elif file_type == 'xlsx':
            wb = openpyxl.load_workbook(filepath, data_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                extracted_text += f"\n--- Sheet: {sheet} ---\n"
                for row in ws.iter_rows(values_only=True):
                    row_text = "\t".join([str(cell) for cell in row if cell is not None])
                    if row_text:
                        extracted_text += row_text + "\n"

        elif file_type == 'pptx':
            prs = pptx.Presentation(filepath)
            for idx, slide in enumerate(prs.slides):
                extracted_text += f"\n--- Slide {idx+1} ---\n"
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        extracted_text += shape.text + "\n"

        elif file_type in ['jpg', 'jpeg', 'png']:
            image = Image.open(filepath)
            reader = load_ocr_reader()
            image_np = np.array(image)
            results = reader.readtext(image_np, detail=0)
            extracted_text = "\n".join(results)

    except Exception as e:
        print(f"Error ekstraksi file: {e}")

    return extracted_text

def add_document_to_vector_db(doc_id, filepath, filename, opd_name):
    text = extract_text_from_filepath(filepath)
    if not text.strip():
        return
    
    # Pemotongan teks (Chunking) 1000 karakter
    chunks = [text[i:i+1000] for i in range(0, len(text), 1000)]
    
    conn = sqlite3.connect(DB_RAG)
    cursor = conn.cursor()
    for chunk in chunks:
        cursor.execute(
            "INSERT INTO document_chunks (doc_id, filename, opd_name, chunk_text) VALUES (?, ?, ?, ?)",
            (doc_id, filename, opd_name, chunk)
        )
    conn.commit()
    conn.close()

def remove_document_from_vector_db(doc_id):
    conn = sqlite3.connect(DB_RAG)
    cursor = conn.cursor()
    cursor.execute("DELETE FROM document_chunks WHERE doc_id = ?", (doc_id,))
    conn.commit()
    conn.close()

def query_relevant_context(user_query, user_role, user_opd):
    conn = sqlite3.connect(DB_RAG)
    cursor = conn.cursor()
    
    # Ambil potongan dokumen berdasarkan hak akses OPD/Ekbang
    if user_role == 'ekbang':
        cursor.execute("SELECT filename, chunk_text FROM document_chunks")
    else:
        cursor.execute("SELECT filename, chunk_text FROM document_chunks WHERE opd_name = ?", (user_opd,))
        
    rows = cursor.fetchall()
    conn.close()
    
    if not rows:
        return ""

    # Pencarian sederhana berdasarkan kata kunci terbanyak yang cocok dengan pertanyaan
    query_words = set(user_query.lower().split())
    scored_chunks = []
    
    for filename, chunk in rows:
        chunk_words = set(chunk.lower().split())
        score = len(query_words.intersection(chunk_words))
        scored_chunks.append((score, filename, chunk))
        
    # Urutkan berdasarkan skor tertinggi
    scored_chunks.sort(key=lambda x: x[0], reverse=True)
    
    # Ambil 5 potongan teks yang paling relevan
    top_chunks = scored_chunks[:5]
    
    context = ""
    for score, filename, chunk in top_chunks:
        context += f"[File: {filename}]\n{chunk}\n---\n"
        
    return context

def generate_sample_chart(chart_type="bar"):
    labels = ['Kategori A', 'Kategori B', 'Kategori C', 'Kategori D']
    values = [40, 25, 70, 55]
    if chart_type == "line":
        fig = px.line(x=labels, y=values, labels={'x':'Program', 'y':'Capaian (%)'}, title="Grafik Capaian Program Kerja")
    else:
        fig = px.bar(x=labels, y=values, labels={'x':'Kategori', 'y':'Nilai'}, title="Grafik Analisis Dokumen Bappeda")
    return fig